import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "bank-agent-demo" / "python"))

from langchain_core.documents import Document

import build_vector_store as builder_module


class BuildVectorStoreTest(unittest.TestCase):
    def test_complex_excel_inherits_vertical_merges_and_skips_group_header(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "qa.xlsx")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "客户经理QA"
            sheet.append(["问答库", None, None, None])
            sheet.append(["问题索引", None, "标准服务口径", None])
            sheet.append(["项目", "业务大类", "标准问题", "标准答案"])
            sheet.append(["QA-001", "客户分层", "金穗级门槛？", "50万元"])
            sheet.append(["QA-002", None, "白金级门槛？", "200万元"])
            sheet.merge_cells("B4:B5")
            workbook.save(path)

            chunks = builder_module.parse_general_excel(str(path))

            self.assertEqual(2, len(chunks))
            self.assertTrue(all("业务大类:客户分层" in chunk.page_content for chunk in chunks))
            self.assertTrue(any("白金级门槛？" in chunk.page_content for chunk in chunks))
            self.assertFalse(any("问题索引" in chunk.page_content for chunk in chunks))

    def test_canonical_builder_handles_pdf_excel_and_markdown(self):
        with tempfile.TemporaryDirectory() as documents_dir, tempfile.TemporaryDirectory() as parent:
            for filename in ("a.pdf", "b.xlsx", "c.md", "ignored.txt"):
                Path(documents_dir, filename).write_text("content", encoding="utf-8")
            target = Path(parent, "new-index")
            calls = []
            store = SimpleNamespace(persist=lambda: calls.append("persist"))

            class FakeChroma:
                @staticmethod
                def from_documents(documents, embedding, persist_directory):
                    calls.append((documents, embedding, persist_directory))
                    return store

            fake_module = SimpleNamespace(Chroma=FakeChroma)

            def fake_load(path):
                return [Document(page_content=Path(path).name, metadata={"source": Path(path).name})]

            with patch.dict(sys.modules, {"langchain_community.vectorstores": fake_module}), patch.object(
                builder_module, "load_document", side_effect=fake_load
            ):
                result = builder_module.build_vector_store(
                    documents_dir=documents_dir,
                    vector_db_dir=str(target),
                    embedding="embedding",
                )

            indexed_names = {document.metadata["source"] for document in calls[0][0]}
            self.assertIs(store, result)
            self.assertEqual({"a.pdf", "b.xlsx", "c.md"}, indexed_names)
            self.assertEqual("embedding", calls[0][1])
            self.assertEqual(str(target), calls[0][2])
            self.assertEqual("persist", calls[1])


if __name__ == "__main__":
    unittest.main()
